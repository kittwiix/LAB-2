# -*- coding: utf-8 -*-
"""Сохранение расчётов в sqlite3 и экспорт в docx/xlsx."""
import sqlite3
from datetime import datetime

from docx import Document
from openpyxl import Workbook


class CalculationsDB:
    """Обёртка над sqlite3: таблица сохранённых расчётов."""

    def __init__(self, path='calculations.db'):
        self._conn = sqlite3.connect(path)
        self._conn.execute(
            'CREATE TABLE IF NOT EXISTS calculations ('
            'id INTEGER PRIMARY KEY AUTOINCREMENT, '
            'material TEXT, area REAL, params TEXT, '
            'units INTEGER, cost REAL, created TEXT)'
        )

    def add(self, material):
        """Сохранить расчёт материала, вернуть id записи."""
        cur = self._conn.execute(
            'INSERT INTO calculations (material, area, params, units, cost, created) '
            'VALUES (?, ?, ?, ?, ?, ?)',
            (material.name, material.area, material.params_str(),
             material.units_to_buy(), round(material.total_cost(), 2),
             datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        )
        self._conn.commit()
        return cur.lastrowid

    def all(self):
        """Все записи таблицы в порядке добавления."""
        return self._conn.execute(
            'SELECT id, material, area, params, units, cost, created '
            'FROM calculations ORDER BY id'
        ).fetchall()

    def clear(self):
        """Удалить все записи."""
        self._conn.execute('DELETE FROM calculations')
        self._conn.commit()

    def close(self):
        """Закрыть соединение с базой."""
        self._conn.close()

    def __len__(self):
        return self._conn.execute('SELECT COUNT(*) FROM calculations').fetchone()[0]

    def __iter__(self):
        return iter(self.all())


def export_docx(rows, filename='report.docx'):
    """Экспорт таблицы расчётов в документ Word."""
    doc = Document()
    doc.add_heading('Расчёт закупки отделочных материалов', level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = 'Table Grid'
    headers = ('ID', 'Материал', 'Площадь, м²', 'Параметры',
               'К закупке', 'Стоимость, руб.', 'Дата')
    for i, title in enumerate(headers):
        table.rows[0].cells[i].text = title
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = f'{value:.2f}' if isinstance(value, float) else str(value)
    doc.save(filename)
    return filename


def export_xlsx(rows, filename='report.xlsx'):
    """Экспорт таблицы расчётов в книгу Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Расчёты'
    ws.append(('ID', 'Материал', 'Площадь, м²', 'Параметры',
               'К закупке', 'Стоимость, руб.', 'Дата'))
    for row in rows:
        ws.append(row)
    for col, width in zip('ABCDEFG', (6, 12, 12, 22, 14, 16, 20)):
        ws.column_dimensions[col].width = width
    wb.save(filename)
    return filename
