# -*- coding: utf-8 -*-
"""Тесты пакета finishing."""
import os
import sys

import pytest
from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from finishing import (CalculationsDB, FinishMaterial, Laminate, Tile,
                       Wallpaper, export_docx, export_xlsx)


def test_wallpaper_without_reserve():
    # 53 / 5.3 — ровно 10 рулонов, погрешность чисел с плавающей точкой не мешает
    wallpaper = Wallpaper(53, 0.53, 10.0, 590, 0)
    assert wallpaper.units_to_buy() == 10


def test_wallpaper_with_reserve():
    # 53 * 1.1 / 5.3 = 11 рулонов
    wallpaper = Wallpaper(53, 0.53, 10.0, 590, 10)
    assert wallpaper.units_to_buy() == 11


def test_wallpaper_cost():
    wallpaper = Wallpaper(53, 0.53, 10.0, 590, 10)
    assert wallpaper.total_cost() == pytest.approx(11 * 590)


def test_wallpaper_incomplete_roll():
    # 11 / 5.3 = 2.07..., нужно 3 рулона
    wallpaper = Wallpaper(10, 0.53, 10.0, 590, 10)
    assert wallpaper.units_to_buy() == 3


def test_tile_units_and_cost():
    # 25 * 1.1 / 0.09 = 305.5..., нужно 306 плиток
    tile = Tile(25, 0.3, 0.3, 45, 10)
    assert tile.units_to_buy() == 306
    assert tile.total_cost() == pytest.approx(306 * 45)


def test_laminate_units_and_cost():
    # 40 * 1.05 / 2.13 = 19.71..., нужно 20 упаковок
    laminate = Laminate(40, 2.13, 899, 5)
    assert laminate.units_to_buy() == 20
    assert laminate.total_cost() == pytest.approx(20 * 899)


def test_work_area():
    assert Tile(20, 0.3, 0.3, 45, 10).work_area() == pytest.approx(22)
    assert Tile(20, 0.3, 0.3, 45, 0).work_area() == pytest.approx(20)


@pytest.mark.parametrize('bad_area', [0, -5, float('nan'), float('inf'), True, '10'])
def test_bad_area(bad_area):
    with pytest.raises(ValueError):
        Wallpaper(bad_area, 0.53, 10.0, 590, 10)


def test_bad_width():
    with pytest.raises(ValueError):
        Wallpaper(10, 0, 10.0, 590, 10)


def test_bad_price():
    with pytest.raises(ValueError):
        Laminate(10, 2.13, -100, 10)


def test_bad_reserve():
    with pytest.raises(ValueError):
        Tile(10, 0.3, 0.3, 45, -1)


def test_setters_validate():
    wallpaper = Wallpaper(10, 0.53, 10.0, 590, 10)
    with pytest.raises(ValueError):
        wallpaper.area = 0
    with pytest.raises(ValueError):
        wallpaper.price = -1


def test_empty_name():
    wallpaper = Wallpaper(10, 0.53, 10.0, 590, 10)
    with pytest.raises(ValueError):
        wallpaper.name = '   '


def test_abstract_class_cannot_be_created():
    with pytest.raises(TypeError):
        FinishMaterial('Обои', 10, 100, 10)


def test_abstract_methods():
    assert set(FinishMaterial.__abstractmethods__) == {'unit_area', 'unit_title', 'params_str'}


def test_str_contains_results():
    text = str(Tile(25, 0.3, 0.3, 45, 10))
    assert 'Плитка' in text
    assert '306' in text
    assert '13770' in text


def test_repr_shows_constructor_params():
    assert repr(Wallpaper(10, 0.53, 10.0, 590, 10)) == 'Wallpaper(10, 0.53, 10, 590, 10)'
    assert repr(Laminate(40, 2.13, 899, 5)) == 'Laminate(40, 2.13, 899, 5)'


def test_eq():
    assert Wallpaper(10, 0.53, 10.0, 590, 10) == Wallpaper(10, 0.53, 10.0, 590, 10)
    assert Wallpaper(10, 0.53, 10.0, 590, 10) != Wallpaper(10, 1.06, 10.0, 590, 10)
    assert Wallpaper(10, 0.53, 10.0, 590, 10) != Tile(10, 0.53, 10.0, 590, 10)
    assert Wallpaper(10, 0.53, 10.0, 590, 10) != 'обои'


def test_db_add_and_read(tmp_path):
    db = CalculationsDB(str(tmp_path / 'test.db'))
    row_id = db.add(Tile(25, 0.3, 0.3, 45, 10))
    rows = db.all()
    assert len(db) == 1
    assert row_id == rows[0][0]
    assert rows[0][1] == 'Плитка'
    assert rows[0][2] == pytest.approx(25)
    assert rows[0][4] == 306
    assert rows[0][5] == pytest.approx(13770)
    db.close()


def test_db_iter_and_clear(tmp_path):
    db = CalculationsDB(str(tmp_path / 'test.db'))
    db.add(Wallpaper(53, 0.53, 10.0, 590, 10))
    db.add(Laminate(40, 2.13, 899, 5))
    assert [row[1] for row in db] == ['Обои', 'Ламинат']
    db.clear()
    assert len(db) == 0
    db.close()


def test_export_docx(tmp_path):
    db = CalculationsDB(str(tmp_path / 'test.db'))
    db.add(Wallpaper(53, 0.53, 10.0, 590, 10))
    path = export_docx(db.all(), str(tmp_path / 'report.docx'))
    db.close()
    doc = Document(path)
    cells = doc.tables[0].rows[1].cells
    assert cells[1].text == 'Обои'
    assert cells[4].text == '11'
    assert cells[5].text == '6490.00'


def test_export_xlsx(tmp_path):
    db = CalculationsDB(str(tmp_path / 'test.db'))
    db.add(Wallpaper(53, 0.53, 10.0, 590, 10))
    path = export_xlsx(db.all(), str(tmp_path / 'report.xlsx'))
    db.close()
    sheet = load_workbook(path).active
    assert sheet['B2'].value == 'Обои'
    assert sheet['E2'].value == 11
    assert sheet['F2'].value == pytest.approx(6490)
